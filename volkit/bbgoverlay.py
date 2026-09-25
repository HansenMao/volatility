"""A Bloomberg sheet the Vol bulk processing screen loads as its overlay.

``files/bbg_overlay.xlsx`` pulls the last price of the ATM, the 25- and 10-delta risk reversal
and the 25- and 10-delta butterfly for every pair and tenor the workbook's ``CONFIG`` names, and
lays them out as an overlay (``overlay.py``): its **first** sheet is ``pair, tenor, atm, rr25,
rr10, bf25, bf10``, one row per pair and tenor.  Two more sheets carry the machinery -- ``bbg``
(the tickers and the ``=BDP()`` pulls) and ``settings`` (the workbook, the pricing source, the
tenor codes).  The desk opens it on a Bloomberg terminal, lets it fill, saves, and loads it.

What makes it loadable, each point a way the overlay reader would otherwise refuse the file:

* **Nothing but a number or a blank reaches the first sheet.**  A pull that has not come back
  reads ``#N/A Requesting Data...`` and a ticker Bloomberg does not have reads ``#N/A Invalid
  Security`` -- both text, and one text cell in a quote column refuses the whole file.  The
  first sheet takes a pull only ``IF(ISNUMBER(...))``, so a missing quote is a blank, and a
  blank falls through to the book.  How many came back is counted on the comment row above the
  header, which the reader skips, so the file says how complete it is.
* **The reader reads saved values.**  The sheet must be saved *after* the pulls arrive.
* **Every row is unique.**  ``CONFIG`` lists pairs and tenors in two columns, and the rows are
  their product, each once.

The list follows the workbook: ``files/volkit_bbg_overlay.bas``, imported once and saved as
``.xlsm``, re-reads ``CONFIG`` every time the sheet is opened (``Auto_Open``) and on demand.  The
macro writes only the pair and tenor cells and fills the formulas below down to match, so the
formulas here (``FORMULAS``) are the ones it writes, and a test holds the two together.

``python -m volkit.bbgoverlay WORKBOOK OUT`` writes the sheet for a workbook's list as it stands.
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

#: The first data row on the ``overlay`` and ``bbg`` sheets: row 1 is the comment, row 2 the header.
FIRST = 3
#: The five quotes, in overlay order, with the Bloomberg infix each is quoted under
#: (``EURUSDV1M``, ``EURUSD25R1M``, ``EURUSD10B1M`` ...).
QUOTES = (("atm", "V"), ("rr25", "25R"), ("rr10", "10R"), ("bf25", "25B"), ("bf10", "10B"))
#: A tenor as ``CONFIG`` may write it, and the code Bloomberg's vol tickers use.  Anything not
#: here is used upper-cased (``1m`` -> ``1M``); the table sits on the settings sheet to be edited.
TENOR_CODES = (("1D", "ON"), ("ON", "ON"), ("O/N", "ON"), ("OVERNIGHT", "ON"),
               ("1WK", "1W"), ("12M", "1Y"), ("24M", "2Y"), ("36M", "3Y"), ("48M", "4Y"),
               ("60M", "5Y"))
#: How far down the formulas are counted and cleared; the list is never near it (25 x 12 = 300).
LAST = 5000

_OV_COLS = "CDEFG"          # overlay: atm rr25 rr10 bf25 bf10
_TK_COLS = "DEFGH"          # bbg: the five tickers
_PX_COLS = "IJKLM"          # bbg: the five pulls

#: The formulas of data row ``FIRST``, relative, by sheet and column.  The macro writes these
#: same strings into the first row of a range and Excel moves them down, as a fill would.
FORMULAS: dict[tuple[str, str], str] = {}
for _i, (_name, _infix) in enumerate(QUOTES):
    FORMULAS[("overlay", _OV_COLS[_i])] = (
        f'=IF(ISNUMBER(bbg!{_PX_COLS[_i]}{FIRST}),bbg!{_PX_COLS[_i]}{FIRST},"")')
    FORMULAS[("bbg", _TK_COLS[_i])] = (
        f'=IF($A{FIRST}="","",$A{FIRST}&"{_infix}"&$C{FIRST}&settings!$B$5&" Curncy")')
    FORMULAS[("bbg", _PX_COLS[_i])] = (
        f'=IF({_TK_COLS[_i]}{FIRST}="","",BDP({_TK_COLS[_i]}{FIRST},settings!$B$4))')
FORMULAS[("bbg", "A")] = f'=IF(overlay!$A{FIRST}="","",overlay!$A{FIRST})'
FORMULAS[("bbg", "B")] = f'=IF(overlay!$B{FIRST}="","",overlay!$B{FIRST})'
FORMULAS[("bbg", "C")] = (
    f'=IF($B{FIRST}="","",IFERROR(VLOOKUP(UPPER(TRIM($B{FIRST})),settings!$D$3:$E$60,2,FALSE),'
    f'UPPER(TRIM($B{FIRST}))))')

#: The comment row above the overlay's header: what came back, from which list, when.  It starts
#: with ``#``, which is what makes the overlay reader pass over it.
COMMENT = ('="# volkit overlay from Bloomberg: "&settings!$B$6&" of "&settings!$B$7&'
           '" quotes in -- save after they arrive. List: "&settings!$B$9&" pairs x "&'
           'settings!$B$10&" tenors from "&settings!$B$2&", "&settings!$B$8')

SETTINGS = (
    ("volkit workbook", "vol_marks.xlsx",
     "whose CONFIG names the pairs and tenors; a bare name is beside this file"),
    ("pricing source", "",
     "blank for the terminal's default, or BGN, CMPN ...: goes into the ticker"),
    ("field", "PX_LAST", "what BDP asks for"),
    ("ticker suffix", '=IF(TRIM(B3)="",""," "&TRIM(B3))', "worked out from the source"),
    ("quotes in", f"=COUNT(overlay!C{FIRST}:G{LAST})", "numbers on the overlay sheet now"),
    ("quotes asked", f'=SUMPRODUCT(--(bbg!D{FIRST}:H{LAST}<>""))', "tickers asked for"),
    ("list read", "", "when the macro last read CONFIG, or why it could not"),
    ("pairs", "", ""),
    ("tenors", "", ""),
)

INSTRUCTIONS = (
    "How to use",
    "1. Open on a Bloomberg terminal with the Excel add-in loaded; the pulls fill in by themselves.",
    "2. When 'quotes in' stops rising, save. volkit reads what was saved, never the live cells.",
    "3. Vol bulk processing > overlay > Browse... this file. A blank quote falls through to the book.",
    "The list: import volkit_bbg_overlay.bas (Alt+F11, File > Import) and save as .xlsm once.",
    "   It then re-reads CONFIG every time this file opens; VolkitBbgRefreshList does it on demand.",
    "Keep 'overlay' the first sheet: the overlay reader reads the first sheet only.",
)


def config_list(workbook) -> tuple[list[str], list[str]]:
    """The pairs and the tenors ``CONFIG`` lists, in its order, each once.

    Read the way ``marketdata`` reads it, and the way the macro does: the header is the first
    row; the pairs are the first of the ``PAIRS`` spellings (the legacy ``BASE`` among them)
    plus a legacy ``COR`` column; the tenors are ``TENORS``.
    """
    from .configsheets import open_workbook
    from .marketdata import CONFIG_PAIR_COLUMNS, _norm
    wb = open_workbook(workbook)
    try:
        if "CONFIG" not in wb.sheetnames:
            raise ValueError(f"{workbook} has no CONFIG sheet")
        rows = [list(r) for r in wb["CONFIG"].iter_rows(values_only=True)]
    finally:
        wb.close()
    if not rows:
        raise ValueError(f"{workbook}: CONFIG is empty")
    head = {_norm(c): i for i, c in reversed(list(enumerate(rows[0]))) if c is not None}
    pcol = next((head[k] for k in CONFIG_PAIR_COLUMNS if k in head), None)
    tcol = head.get("tenors")
    if pcol is None or tcol is None:
        raise ValueError(f"{workbook}: CONFIG needs a PAIRS and a TENORS column; its header "
                         f"is {[c for c in rows[0] if c is not None]}")
    pairs, tenors = [], []
    for col in (pcol, head.get("cor")):
        if col is None:
            continue
        for r in rows[1:]:
            v = str(r[col]).strip().upper() if col < len(r) and r[col] is not None else ""
            if len(v) == 6 and v.isalpha() and v not in pairs:
                pairs.append(v)
    for r in rows[1:]:
        v = str(r[tcol]).strip() if tcol < len(r) and r[tcol] is not None else ""
        if v and v.upper() not in [t.upper() for t in tenors]:
            tenors.append(v)
    return pairs, tenors


def write(workbook, out, *, source_name: str | None = None) -> dict:
    """Write the sheet for ``workbook``'s list as it stands.  Returns what it wrote."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    pairs, tenors = config_list(workbook)
    rows = [(p, t) for p in pairs for t in tenors]
    wb = openpyxl.Workbook()
    ov = wb.active
    ov.title = "overlay"
    bb = wb.create_sheet("bbg")
    st = wb.create_sheet("settings")
    bold = Font(bold=True)
    grey = Font(italic=True, color="666666")

    ov["A1"] = COMMENT
    ov["A1"].font = grey
    for j, h in enumerate(["pair", "tenor"] + [q for q, _ in QUOTES]):
        c = ov.cell(row=2, column=j + 1, value=h)
        c.font = bold
    bb["A1"] = "# the pulls behind the overlay sheet; the macro keeps these rows in step with it"
    bb["A1"].font = grey
    for j, h in enumerate(["pair", "tenor", "bbg tenor"] + [f"{q} ticker" for q, _ in QUOTES]
                          + [f"{q} {'PX_LAST'}" for q, _ in QUOTES]):
        bb.cell(row=2, column=j + 1, value=h).font = bold

    for n, (pair, tenor) in enumerate(rows):
        r = FIRST + n
        ov.cell(row=r, column=1, value=pair)
        ov.cell(row=r, column=2, value=tenor)
        for (sheet, col), f in FORMULAS.items():
            ws = ov if sheet == "overlay" else bb
            ws[f"{col}{r}"] = _shift(f, r)

    st["A1"] = "setting"
    st["B1"] = "value"
    st["C1"] = "what it is"
    for c in ("A1", "B1", "C1"):
        st[c].font = bold
    for i, (name, value, what) in enumerate(SETTINGS, start=2):
        st.cell(row=i, column=1, value=name)
        st.cell(row=i, column=2, value=value)
        st.cell(row=i, column=3, value=what).font = grey
    st["B2"] = source_name or Path(workbook).name
    st["B8"] = f"written by volkit {datetime.now():%Y-%m-%d %H:%M} from {Path(workbook).name}"
    st["B9"] = len(pairs)
    st["B10"] = len(tenors)
    for c in ("B2", "B3", "B4"):
        st[c].fill = PatternFill("solid", fgColor="FFF2CC")
    st["D1"] = "tenor as CONFIG writes it"
    st["E1"] = "Bloomberg code"
    st["D1"].font = st["E1"].font = bold
    st["D2"] = "(anything not listed is used upper-cased: 1m -> 1M)"
    st["D2"].font = grey
    for i, (a, b) in enumerate(TENOR_CODES, start=3):
        st.cell(row=i, column=4, value=a)
        st.cell(row=i, column=5, value=b)
    for i, line in enumerate(INSTRUCTIONS, start=1):
        st.cell(row=i, column=7, value=line).font = bold if i == 1 else Font()

    for ws, widths in ((ov, (9, 7, 8, 8, 8, 8, 8)),
                       (bb, (9, 7, 9) + (22,) * 5 + (10,) * 5),
                       (st, (16, 22, 58, 26, 16, 2, 90))):
        for j, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j + 1)].width = w
    ov.freeze_panes = "A3"
    bb.freeze_panes = "A3"
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return {"pairs": pairs, "tenors": tenors, "rows": len(rows), "path": str(out)}


def _shift(formula: str, row: int) -> str:
    """The row-``FIRST`` formula moved to ``row``: every relative ``<col><FIRST>`` reference.

    Only a reference to the data row is relative; ``settings!$B$4`` and the lookup range are
    absolute and stay put, which is what a fill down does too.
    """
    import re
    return re.sub(rf"(?<![$\w])(\$?[A-M]){FIRST}(?!\d)", lambda m: f"{m.group(1)}{row}", formula)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("usage: python -m volkit.bbgoverlay WORKBOOK OUT.xlsx", file=sys.stderr)
        raise SystemExit(2)
    got = write(sys.argv[1], sys.argv[2])
    print(f"wrote {got['path']}: {len(got['pairs'])} pairs x {len(got['tenors'])} tenors = "
          f"{got['rows']} rows")
