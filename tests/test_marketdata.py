"""The feed, the market boxes and cross levels built from the legs.

Split out of the old 15,008-line ``tests/test_volkit.py``; the shared imports,
paths and helpers are in ``tests/_support.py``.
"""

from __future__ import annotations

from ._support import *  # noqa: F401,F403


class TestMarketData(unittest.TestCase):
    def _workbook(self, config: dict, pairs: list) -> Path:
        """A minimal workbook with the given CONFIG columns, in a temp dir.

        PARAMS carries whatever ``pairs`` names, so a test can say what the
        sheet declares and nothing else.  Every pair also gets a smile sheet,
        because a real workbook has one: a pair CONFIG names with no sheet
        behind it is now a reported problem rather than a silent skip, and a
        fixture that left them out would be asserting on a workbook nobody
        would ship.
        """
        import tempfile
        d = Path(tempfile.mkdtemp())
        path = d / "book.xlsx"
        self.addCleanup(shutil.rmtree, d, True)
        params = pd.DataFrame(
            {p: [8.0, 9.0, 0.0, 0.0, 5.0, 0.0, 50.0] for p in pairs},
            index=["initial", "long term", "ratevol", "addon", "MR",
                   "rate corr", "short decay"],
        )
        for p in pairs:
            if "USD" not in (p[:3], p[3:6]):
                params[p] = [0.4, 0.3, 0.0, 0.0, 4.0, 0.0, 50.0]
        tenors = [t for t in config.get("TENORS", ["1m", "3m"]) if t]
        smile = pd.DataFrame({
            "expiry": tenors,
            "ST 10D": [0.60] * len(tenors),
            "ST 25D": [0.20] * len(tenors),
            "RR 25D": [-0.10] * len(tenors),
            "RR 10D": [-0.19] * len(tenors),
        })
        with pd.ExcelWriter(path) as xw:
            width = max(len(v) for v in config.values())
            padded = {k: list(v) + [None] * (width - len(v)) for k, v in config.items()}
            pd.DataFrame(padded).to_excel(xw, sheet_name="CONFIG", index=False)
            params.to_excel(xw, sheet_name="PARAMS")
            for p in pairs:
                smile.to_excel(xw, sheet_name=p, index=False)
        return path

    def test_a_currency_column_on_the_events_sheet_weights_every_pair_with_that_leg(self):
        """The EVENTS sheet: a column headed USD is the dollar's weight on
        each row and is shared by every pair with a dollar leg; a pair's cell
        on that row is its adjustment on top.  Events used to be dated rows
        on PARAMS, which gave one weight two homes -- the old bug this pins
        is a weight typed on USDJPY and not on EURUSD, and the two pairs
        disagreeing about what the Fed is worth."""
        import tempfile
        d = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, d, True)
        path = d / "book.xlsx"
        idx = ["initial", "long term", "ratevol", "addon", "MR", "rate corr", "short decay"]
        params = pd.DataFrame({
            "USDJPY": [8.0, 9.0, 0, 0, 5.0, 0, 50.0],
            "EURUSD": [8.0, 9.0, 0, 0, 5.0, 0, 50.0],
        }, index=idx)
        # The sheet is typed in Hong Kong time, like every event in this tool.
        events = pd.DataFrame({
            None: ["2026-09-17 06:00", "2026-10-30 19:00"],
            "USD": [1.5, 0.0],
            "JPY": [0.3, 1.5],
            "USDJPY": [0.2, None],
            "EURUSD": [None, None],
        })
        smile = pd.DataFrame({"expiry": ["1m", "3m"], "ST 10D": [0.6, 0.6],
                              "ST 25D": [0.2, 0.2], "RR 25D": [-0.1, -0.1],
                              "RR 10D": [-0.19, -0.19]})
        with pd.ExcelWriter(path) as xw:
            pd.DataFrame({"PAIRS": ["USDJPY", "EURUSD"], "TENORS": ["1m", "3m"]}).to_excel(
                xw, sheet_name="CONFIG", index=False)
            params.to_excel(xw, sheet_name="PARAMS")
            events.to_excel(xw, sheet_name="EVENTS", index=False)
            for name in ("USDJPY", "EURUSD"):
                smile.to_excel(xw, sheet_name=name, index=False)
        data = ExcelSource(path).load()
        self.assertEqual(data.problems, [], data.problems)
        # 06:00 Hong Kong is 22:00 UTC the day before: the sheet's clock is
        # the workbook's, and the model's is UTC.
        uj = {e.when.strftime("%d%b"): e for e in data.events.for_pair("USDJPY")}
        self.assertAlmostEqual(uj["16Sep"].bump, 0.020)          # 1.5 + 0.3 + 0.2
        self.assertEqual(uj["16Sep"].weights, {"USD": 0.015, "JPY": 0.003})
        self.assertAlmostEqual(uj["16Sep"].adjust, 0.002)
        self.assertAlmostEqual(uj["30Oct"].bump, 0.015)          # the JPY leg alone
        eu = {e.when.strftime("%d%b"): e for e in data.events.for_pair("EURUSD")}
        self.assertAlmostEqual(eu["16Sep"].bump, 0.015)          # no JPY leg, no cell
        self.assertAlmostEqual(eu["30Oct"].bump, 0.0)            # nothing on either leg
        # The curve only takes the rows that move it; the panel sees them all.
        touching = data.events.for_pair("EURUSD", touching_only=True)
        self.assertEqual([e.when.strftime("%d%b") for e in touching], ["16Sep"])
        self.assertTrue(any("EVENTS: 2 event(s)" in n for n in data.notes), data.notes)

    def test_a_date_row_left_on_params_is_reported_not_read(self):
        """Events had dated rows on PARAMS once.  One left behind must not be
        read -- two homes for one bump is how a weight comes to mean two
        things -- and must not be ignored either."""
        import tempfile
        d = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, d, True)
        path = d / "book.xlsx"
        idx = ["initial", "long term", "ratevol", "addon", "MR", "rate corr", "short decay",
               "2026-09-16 22:00"]
        params = pd.DataFrame({"USDJPY": [8.0, 9.0, 0, 0, 5.0, 0, 50.0, 0.2]}, index=idx)
        smile = pd.DataFrame({"expiry": ["1m"], "ST 10D": [0.6], "ST 25D": [0.2],
                              "RR 25D": [-0.1], "RR 10D": [-0.19]})
        with pd.ExcelWriter(path) as xw:
            pd.DataFrame({"PAIRS": ["USDJPY"], "TENORS": ["1m"]}).to_excel(
                xw, sheet_name="CONFIG", index=False)
            params.to_excel(xw, sheet_name="PARAMS")
            smile.to_excel(xw, sheet_name="USDJPY", index=False)
        data = ExcelSource(path).load()
        self.assertTrue(any("EVENTS" in p and "is a date" in p for p in data.problems),
                        data.problems)
        self.assertEqual(data.events.rows, [])

    def test_an_events_column_that_is_neither_a_currency_nor_a_pair_is_reported(self):
        """A cell nothing reads is a bump nobody gets."""
        import tempfile
        d = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, d, True)
        path = d / "book.xlsx"
        idx = ["initial", "long term", "ratevol", "addon", "MR", "rate corr", "short decay"]
        params = pd.DataFrame({"USDJPY": [8.0, 9.0, 0, 0, 5.0, 0, 50.0]}, index=idx)
        events = pd.DataFrame({None: ["2026-09-17 06:00"], "USD": [1.5], "USDNOK": [2.0]})
        smile = pd.DataFrame({"expiry": ["1m"], "ST 10D": [0.6], "ST 25D": [0.2],
                              "RR 25D": [-0.1], "RR 10D": [-0.19]})
        with pd.ExcelWriter(path) as xw:
            pd.DataFrame({"PAIRS": ["USDJPY"], "TENORS": ["1m"]}).to_excel(
                xw, sheet_name="CONFIG", index=False)
            params.to_excel(xw, sheet_name="PARAMS")
            events.to_excel(xw, sheet_name="EVENTS", index=False)
            smile.to_excel(xw, sheet_name="USDJPY", index=False)
        data = ExcelSource(path).load()
        self.assertTrue(any("USDNOK" in p for p in data.problems), data.problems)

    def test_real_workbook_loads_without_problems(self):
        data = ExcelSource(BOOK).load()
        self.assertEqual(data.problems, [], data.problems)
        self.assertIn("USDJPY", data.pairs)
        self.assertTrue(data.pairs["AUDJPY"].is_cross)
        self.assertEqual(data.pairs["AUDJPY"].legs, ("AUDUSD", "USDJPY"))

    def test_the_shipped_config_is_two_columns(self):
        """The sheet a desk maintains is pairs and tenors, and nothing else.

        The COR column and the column-per-cross naming its legs are gone: a
        cross has exactly one sensible pair of dollar legs, so it is worked
        out from the name instead of written down thirteen times.
        """
        with marketdata.open_workbook(BOOK) as xls:
            cfg = pd.read_excel(xls, "CONFIG")
        self.assertEqual([str(c).strip().upper() for c in cfg.columns],
                         ["PAIRS", "TENORS"])
        data = ExcelSource(BOOK).load()
        self.assertTrue(data.pairs["EURGBP"].derived)
        self.assertEqual(data.pairs["EURGBP"].legs, ("EURUSD", "GBPUSD"))

    def test_both_layouts_of_the_shipped_workbook_agree(self):
        """The same marks under both CONFIG layouts are the same book.

        ``vol_marks_legacy_format.xlsx`` is the sheet as it was -- BASE, COR
        and a column per cross -- kept because the legacy tool in the repo
        root reads only that.  It is also the strongest guard there is on the
        derivation: the legs it names by hand are the legs the two-column
        sheet works out from the names.
        """
        legacy = WORKBOOK.parent / "vol_marks_legacy_format.xlsx"
        new, old = ExcelSource(WORKBOOK).load(), ExcelSource(legacy).load()
        self.assertEqual(old.problems, [], old.problems)
        self.assertEqual(sorted(new.pairs), sorted(old.pairs))
        for name, spec in new.pairs.items():
            self.assertEqual(spec.is_cross, old.pairs[name].is_cross, name)
            self.assertEqual(spec.legs, old.pairs[name].legs, name)
        self.assertEqual(new.tenor_points, old.tenor_points)

    def test_a_cross_is_broken_into_two_dollar_pairs(self):
        path = self._workbook({"PAIRS": ["EURJPY"], "TENORS": ["1m", "3m"]},
                              ["EURJPY", "EURUSD", "USDJPY"])
        data = ExcelSource(path).load()
        self.assertEqual(data.problems, [], data.problems)
        self.assertTrue(data.pairs["EURJPY"].is_cross)
        self.assertTrue(data.pairs["EURJPY"].derived)
        self.assertEqual(data.pairs["EURJPY"].legs, ("EURUSD", "USDJPY"))
        # The legs nobody listed are there, and say which cross wanted them.
        self.assertEqual(data.pairs["EURUSD"].implied_by, "EURJPY")
        self.assertEqual(data.pairs["USDJPY"].implied_by, "EURJPY")
        self.assertFalse(data.pairs["EURUSD"].is_cross)
        self.assertTrue(any("EURJPY" in n for n in data.notes), data.notes)

    def test_a_derived_cross_is_marked_by_correlation(self):
        """The mechanism is the correlation, not a backbone of its own.

        A cross's initial / long term / MR cells have always meant
        correlation initial / final / decay; deriving the legs must not
        change which of the two a pair gets.
        """
        path = self._workbook({"PAIRS": ["EURJPY"], "TENORS": ["1m", "3m"]},
                              ["EURJPY", "EURUSD", "USDJPY"])
        book = Book.from_excel(path, ASOF).build()
        self.assertIsInstance(book["EURJPY"].atm, CrossAtmCurve)
        # 0.4 -> 0.3 with decay 4: those cells are read as a correlation and
        # are not divided by 100 the way a volatility is.
        self.assertAlmostEqual(float(book["EURJPY"].atm.correlation(0.0)), 0.4)
        self.assertAlmostEqual(float(book["EURJPY"].atm.correlation(50.0)), 0.3)

    def test_a_derived_cross_gets_the_same_legs_the_sheet_used_to_name(self):
        """The orientation is the whole of MIGRATION.md's first entry.

        AUDJPY is AUDUSD and USDJPY -- the dollar in opposite places, so the
        triangle takes +2*rho -- while EURGBP is EURUSD and GBPUSD, the
        dollar in the same place and -2*rho.  A leg written upside down would
        flip that sign silently.
        """
        for pair, legs in (("AUDJPY", ("AUDUSD", "USDJPY")),
                           ("EURGBP", ("EURUSD", "GBPUSD")),
                           ("GBPNZD", ("GBPUSD", "NZDUSD")),
                           ("EURCNH", ("EURUSD", "USDCNH")),
                           ("CNHHKD", ("USDCNH", "USDHKD"))):
            self.assertEqual(dollar_legs(pair), legs)
        self.assertEqual(infer_leg_signs("AUDJPY", *dollar_legs("AUDJPY")), (1, -1))
        self.assertEqual(infer_leg_signs("EURGBP", *dollar_legs("EURGBP")), (1, 1))

    def test_a_pair_config_names_with_no_sheet_is_reported(self):
        """The shipped workbook lost its EURGBP tab to a USDHKD one while
        CONFIG went on naming EURGBP.  Nothing said so: the reader skipped the
        pair in silence, ``volkit check`` reported "no problems found", and
        the first thing to ask that surface for a smile raised "EURGBP: no
        smile term structure; run calibrate() first" -- which names neither
        the workbook nor the tab somebody deleted, and which took the Windows
        build down at the test suite.
        """
        path = self._workbook({"PAIRS": ["USDJPY", "EURUSD"], "TENORS": ["1m", "3m"]},
                              ["USDJPY", "EURUSD"])
        import openpyxl
        wb = openpyxl.load_workbook(path)
        del wb["EURUSD"]
        wb.save(path)
        data = ExcelSource(path).load()
        self.assertEqual(len(data.problems), 1, data.problems)
        self.assertIn("EURUSD", data.problems[0])
        self.assertIn("no 'EURUSD' sheet", data.problems[0])
        self.assertIn("USDJPY", data.marks)
        self.assertNotIn("EURUSD", data.marks)

    def test_a_sheet_whose_rows_cannot_be_read_is_reported(self):
        """The same failure by the other route: the rows are there and unreadable."""
        path = self._workbook({"PAIRS": ["USDJPY"], "TENORS": ["1m", "3m"]}, ["USDJPY"])
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb["USDJPY"]
        for r in range(2, ws.max_row + 1):
            for c in range(2, 6):
                ws.cell(row=r, column=c).value = None
        wb.save(path)
        data = ExcelSource(path).load()
        self.assertTrue(any("no readable quotes" in p for p in data.problems), data.problems)

    def test_a_sheet_with_its_columns_and_no_rows_is_a_pair_not_yet_quoted(self):
        """Which is a real state now that a pair is created from the screens.

        It used to be a problem, because the only way to get here was deleting
        rows by hand.  A pair added in the Config window arrives exactly like
        this, and a check that goes red on a pair somebody has just made is a
        check people stop reading.  Still *said* -- a pair with no smile is
        worth knowing about -- just not called a fault in the workbook.
        """
        path = self._workbook({"PAIRS": ["USDJPY"], "TENORS": ["1m", "3m"]}, ["USDJPY"])
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb["USDJPY"]
        ws.delete_rows(2, ws.max_row)
        wb.save(path)
        data = ExcelSource(path).load()
        self.assertEqual(data.problems, [])
        self.assertTrue(any("no quotes yet" in n for n in data.notes), data.notes)

    def test_a_pair_asked_for_with_no_quotes_says_so(self):
        """``calibrate_smiles`` used to skip it silently, so the book came
        back looking loaded and refused on the first smile."""
        path = self._workbook({"PAIRS": ["USDJPY", "EURUSD"], "TENORS": ["1m", "3m"]},
                              ["USDJPY", "EURUSD"])
        import openpyxl
        wb = openpyxl.load_workbook(path)
        del wb["EURUSD"]
        wb.save(path)
        book = Book.from_excel(path, ASOF).load_all(["USDJPY", "EURUSD"])
        self.assertTrue(any("EURUSD" in w and "no smile quotes" in w for w in book.warnings),
                        book.warnings)

    def test_a_dollar_pair_has_no_legs_to_derive(self):
        with self.assertRaises(ValueError):
            dollar_legs("USDJPY")
        with self.assertRaises(ValueError):
            dollar_legs("EURUR")

    def test_the_legacy_layout_still_loads_and_its_legs_win(self):
        """BASE / COR / a column per cross is still a workbook we read.

        And a sheet that names legs is not second-guessed: the derived legs
        would be EURUSD and USDJPY, and this one says to go through sterling
        instead.  A leg that is itself a cross is then broken down in its
        turn, which is why the legs are resolved on a work list.
        """
        path = self._workbook(
            {"BASE": ["EURUSD", "USDJPY", "GBPUSD"], "COR": ["EURJPY"],
             "EURJPY": ["EURGBP", "GBPJPY"], "TENORS": ["1m", "3m"]},
            ["EURUSD", "USDJPY", "GBPUSD", "EURJPY", "EURGBP", "GBPJPY"])
        data = ExcelSource(path).load()
        self.assertEqual(data.problems, [], data.problems)
        self.assertTrue(data.pairs["EURJPY"].is_cross)
        self.assertFalse(data.pairs["EURJPY"].derived)
        self.assertEqual(data.pairs["EURJPY"].legs, ("EURGBP", "GBPJPY"))
        # GBPJPY was named by nothing but that column, and is a cross itself.
        self.assertEqual(data.pairs["GBPJPY"].implied_by, "EURJPY")
        self.assertEqual(data.pairs["GBPJPY"].legs, ("GBPUSD", "USDJPY"))

    def test_legs_that_cannot_build_the_cross_are_reported_not_raised(self):
        path = self._workbook(
            {"BASE": ["EURUSD", "USDJPY"], "COR": ["EURJPY"],
             "EURJPY": ["EURUSD", "EURUSD"], "TENORS": ["1m"]},
            ["EURUSD", "USDJPY", "EURJPY"])
        data = ExcelSource(path).load()
        self.assertTrue(any("EURJPY" in p for p in data.problems), data.problems)
        self.assertNotIn("EURJPY", data.pairs)

    def test_a_row_that_is_not_a_pair_is_named_and_the_rest_still_load(self):
        path = self._workbook({"PAIRS": ["EURUSD", "EUR/USD", "USDJPY"],
                               "TENORS": ["1m"]},
                              ["EURUSD", "USDJPY"])
        data = ExcelSource(path).load()
        self.assertIn("EURUSD", data.pairs)
        self.assertIn("USDJPY", data.pairs)
        self.assertTrue(any("EUR/USD" in p for p in data.problems), data.problems)

    def test_a_derivation_reaches_the_page(self):
        """A pair that came out of a convention must not read like one that
        was written down, so the note travels with the state and the page
        shows it -- on the meta line's tooltip, not in the message box, which
        holds errors and warnings only."""
        from volkit.webapp import BookService
        state = BookService(str(BOOK), ASOF).state()
        self.assertTrue(any("EURGBP = EURUSD x GBPUSD" in n for n in state["notes"]),
                        state["notes"])
        page = _source("volkit", "web", "index.html")
        self.assertIn("STATE.notes", page)
        self.assertNotIn("nts.map(x=>`<li", page)

    def test_a_config_with_no_pairs_column_says_what_it_wants(self):
        path = self._workbook({"THINGS": ["EURUSD"], "TENORS": ["1m"]}, ["EURUSD"])
        with self.assertRaises(MarketDataError) as ctx:
            ExcelSource(path).load()
        self.assertIn("PAIRS", str(ctx.exception))

    def test_units_are_converted_once(self):
        data = ExcelSource(BOOK).load()
        self.assertAlmostEqual(data.params["USDJPY"].initial, 0.0605)
        self.assertAlmostEqual(data.marks["USDJPY"][0].st_25, 0.002175)

    def test_cross_correlations_are_not_rescaled(self):
        data = ExcelSource(BOOK).load()
        self.assertAlmostEqual(data.params["AUDJPY"].initial, 0.37)

    def test_missing_file_raises_clearly(self):
        with self.assertRaises(MarketDataError):
            ExcelSource("/nonexistent/nope.xlsx")

    def test_premium_adjustment_follows_the_pair(self):
        data = ExcelSource(BOOK).load()
        self.assertTrue(data.pairs["USDJPY"].resolved_premium_adjusted())
        self.assertFalse(data.pairs["EURUSD"].resolved_premium_adjusted())


class TestFeed(unittest.TestCase):
    def setUp(self):
        self.feed = MarketFeed.load(FEED)

    def test_sample_feed_loads_clean(self):
        self.assertEqual(self.feed.problems, [], self.feed.problems)
        self.assertIn("USDJPY", self.feed)

    def test_pip_divisor_by_term_currency(self):
        self.assertEqual(pip_divisor("USDJPY"), 100.0)
        self.assertEqual(pip_divisor("EURUSD"), 10000.0)

    def test_interpolation_is_exact_at_the_pillars(self):
        """Every pillar reads back exactly where the curve puts it.

        Where it puts it is its own **delivery date** -- a broker's 1M swap
        points are the points to the 1M value date, and a month is 30 or 31
        days and not a nominal 30.44.  Asking at ``tenor_to_years`` used to be
        the same question and no longer is; it is now a point a day or so away
        from the pillar, which is exactly the gap this placement closes.
        """
        pf = self.feed.pairs["USDJPY"]
        self.assertIsNotNone(pf.times)
        for t, pts in zip(pf.times, pf.points):
            got, _ = pf.forward_points(t)
            self.assertAlmostEqual(got, pts, places=10)

    def test_a_pillar_is_placed_on_its_own_delivery_date(self):
        """The axis is years from the spot date, and a tenor pillar is on it.

        Pinned because it is what makes an option's forward land *on* a
        pillar rather than between two of them: the option is read at its own
        settlement date, and the pillar is quoted to the same date.
        """
        from volkit.calendars import DEFAULT_CALENDARS
        pf = self.feed.pairs["USDJPY"]
        for tenor, t in zip(pf.tenors, pf.times):
            delivery = DEFAULT_CALENDARS.delivery_date("USDJPY", tenor, self.feed.today)
            self.assertAlmostEqual(
                t, (delivery - pf.spot_date).days / DAYS_IN_YEAR, places=12, msg=tenor)
        # 1M is 31 days here and not the 30.44 a nominal year fraction gives
        self.assertNotAlmostEqual(pf.times[pf.tenors.index("1M")],
                                  tenor_to_years("1M"), places=4)

    def test_interpolation_between_pillars_is_bracketed(self):
        pf = self.feed.pairs["USDJPY"]
        lo, _ = pf.forward_points(tenor_to_years("1m"))
        hi, _ = pf.forward_points(tenor_to_years("2m"))
        mid, extrap = pf.forward_points(0.5 * (tenor_to_years("1m") + tenor_to_years("2m")))
        self.assertFalse(extrap)
        self.assertTrue(min(lo, hi) < mid < max(lo, hi))

    def test_points_go_to_zero_at_the_very_front(self):
        pts, _ = self.feed.pairs["USDJPY"].forward_points(0.0)
        self.assertAlmostEqual(pts, 0.0, places=12)

    def test_beyond_the_last_pillar_is_flagged_not_trended(self):
        pf = self.feed.pairs["USDJPY"]
        far, extrap = pf.forward_points(5.0)
        last, _ = pf.forward_points(tenor_to_years("1y"))
        self.assertTrue(extrap)
        self.assertAlmostEqual(far, last, places=10)

    def test_quote_is_json_safe(self):
        import json
        q = self.feed.quote("USDJPY", 0.2)
        self.assertIsInstance(q["extrapolated"], bool)
        json.dumps(q)   # must not need a default= coercion

    def test_missing_file_and_unknown_pair_raise(self):
        with self.assertRaises(FeedError):
            MarketFeed.load("/nonexistent/feed.csv")
        with self.assertRaises(FeedError):
            self.feed.quote("XXXYYY", 0.2)


class TestDatedFeed(unittest.TestCase):
    """A feed line may name a date instead of a tenor.

    Added 2026-08-31.  The front of a bank's own forward file is not on
    standard tenors: it is the overnight and the tom-next, each quoted as one
    *day* of points rather than as points from spot, and neither of them has a
    tenor to be written as.  Reading them needs a spot date, which is why the
    valuation date is threaded in from the caller's clock rather than taken
    from the machine.
    """

    def feed(self, body, **kw):
        import tempfile
        tmp = tempfile.mkdtemp()
        path = Path(tmp) / "dated.csv"
        path.write_text(textwrap.dedent(body).lstrip(), encoding="utf-8")
        self.addCleanup(shutil.rmtree, tmp, True)
        return MarketFeed.load(path, **kw)

    def test_a_dated_pillar_lands_where_the_same_tenor_does(self):
        """One axis for both spellings, so a mixed file has no seam in it.

        A date 28 days after spot and the tenor a desk would call it are the
        same point on the curve, not two.
        """
        f = self.feed("""
            USDJPY,SPOT,150.25
            USDJPY,SPOT DATE,2026-09-01
            USDJPY,2026-09-29,-11.40
        """)
        self.assertEqual(f.problems, [], f.problems)
        pf = f.pairs["USDJPY"]
        self.assertEqual(pf.tenors, ["2026-09-29"])
        self.assertAlmostEqual(pf.forward_points(28 / DAYS_IN_YEAR)[0], -11.40, places=10)
        # and read by the date it was quoted at, which is the same knot
        self.assertAlmostEqual(pf.points_on("2026-09-29")[0], -11.40, places=10)
        self.assertAlmostEqual(pf.forward_on("2026-09-29")[0], 150.25 - 0.1140, places=10)

    def test_a_date_on_or_before_spot_is_one_day_of_points(self):
        """The near side is stated as rates and accumulated back from spot.

        T/N is the day ending on the spot date, so it prices the day *before*
        spot; O/N the one before that.  Cumulative points at the spot date are
        zero by definition, and everything on the near side is negative of the
        sum of the days between.
        """
        f = self.feed("""
            # asof: 2026-08-27
            USDJPY,SPOT,150.25
            USDJPY,SPOT DATE,2026-08-31
            USDJPY,2026-08-31,-0.40
            USDJPY,2026-08-28,-0.10
        """)
        self.assertEqual(f.problems, [], f.problems)
        pf = f.pairs["USDJPY"]
        self.assertEqual([d.isoformat() for d, _ in pf.daily],
                         ["2026-08-28", "2026-08-31"])
        self.assertEqual(pf.tenors, [])          # nothing on the far side
        self.assertAlmostEqual(pf.points_on("2026-08-31")[0], 0.0, places=12)
        # the day ending on spot is the tom-next: it prices spot minus one
        self.assertAlmostEqual(pf.points_on("2026-08-30")[0], +0.40, places=12)
        # 29-Aug and 30-Aug are unquoted, so their rate is interpolated
        # between -0.10 (28th) and -0.40 (31st): -0.20 and -0.30.
        self.assertAlmostEqual(pf.points_on("2026-08-29")[0], 0.40 + 0.30, places=12)
        self.assertAlmostEqual(pf.points_on("2026-08-28")[0], 0.70 + 0.20, places=12)
        self.assertAlmostEqual(pf.points_on("2026-08-27")[0], 0.90 + 0.10, places=12)
        # and a negative swap point means the earlier date is the higher rate
        self.assertGreater(pf.forward_on("2026-08-27")[0], pf.spot)

    def test_the_near_side_interpolates_rates_and_not_the_running_total(self):
        """The rows are rates, so it is the rates that are interpolated.

        A straight line through the *cumulative* knots at the 28th and the
        31st would put the 30th at +0.30.  It is +0.40, because the day being
        skipped over is the expensive one.
        """
        f = self.feed("""
            # asof: 2026-08-27
            EURUSD,SPOT,1.0842
            EURUSD,SPOT DATE,2026-08-31
            EURUSD,2026-08-31,-0.40
            EURUSD,2026-08-28,-0.10
        """)
        pf = f.pairs["EURUSD"]
        straight_line = 0.90 * (31 - 30) / (31 - 28)
        self.assertNotAlmostEqual(pf.points_on("2026-08-30")[0], straight_line, places=6)
        self.assertAlmostEqual(pf.points_on("2026-08-30")[0], 0.40, places=12)

    def test_a_date_already_delivered_is_passed_over_and_said(self):
        """A forward that has already delivered is not a forward.

        Passed over rather than refused: a published file carrying yesterday's
        row is an ordinary thing, and it is a note and not a problem.  What it
        may never be is silently placed, which would put a knot behind the
        valuation date and drag the front of the curve onto it.
        """
        f = self.feed("""
            # asof: 2026-08-27
            USDJPY,SPOT,150.25
            USDJPY,2026-08-26,-0.10
            USDJPY,2026-09-30,-11.40
        """)
        self.assertEqual(f.problems, [], f.problems)
        self.assertTrue(any("2026-08-26" in n and "passed over" in n for n in f.notes), f.notes)
        pf = f.pairs["USDJPY"]
        self.assertEqual(pf.daily, [])
        self.assertEqual(pf.tenors, ["2026-09-30"])

    def test_a_dated_row_with_no_valuation_date_anywhere_is_refused(self):
        """It is never guessed at.

        The spot date places every dated row, and a spot date needs a day to
        count from.  A wall-clock reading here would be the one call to
        ``utcnow`` inside the model, and would move the whole near side of the
        curve on a valuation in the past.
        """
        f = self.feed("""
            USDJPY,SPOT,150.25
            USDJPY,2026-09-30,-11.40
        """)
        self.assertTrue(any("valuation date" in p for p in f.problems), f.problems)
        self.assertEqual(f.pairs["USDJPY"].tenors, [])
        # and given one, the same file reads
        f2 = self.feed("""
            USDJPY,SPOT,150.25
            USDJPY,2026-09-30,-11.40
        """, today=date(2026, 8, 27))
        self.assertEqual(f2.problems, [], f2.problems)
        self.assertEqual(f2.pairs["USDJPY"].tenors, ["2026-09-30"])

    def test_a_stated_spot_date_beats_the_calendar_and_says_so(self):
        """A publisher knows its own holidays; this tool's calendar may not.

        The near side is placed against that date, so a day's disagreement
        moves the tom-next onto the overnight.  Stated, it is used and said;
        derived, that is said too, because the two must not read alike.
        """
        derived = self.feed("""
            # asof: 2026-08-27
            USDJPY,SPOT,150.25
            USDJPY,2026-09-30,-11.40
        """)
        self.assertEqual(derived.pairs["USDJPY"].spot_date, date(2026, 8, 31))
        self.assertTrue(any("business days after" in n for n in derived.notes), derived.notes)

        stated = self.feed("""
            # asof: 2026-08-27
            USDJPY,SPOT,150.25
            USDJPY,SPOT DATE,2026-09-01
            USDJPY,2026-09-30,-11.40
        """)
        self.assertEqual(stated.pairs["USDJPY"].spot_date, date(2026, 9, 1))
        self.assertTrue(any("as the file states it" in n for n in stated.notes), stated.notes)
        # one day of spot date is one day of curve: the same pillar sits a
        # day further out, so every point interpolated inside it moves.
        self.assertNotAlmostEqual(derived.pairs["USDJPY"].forward_points(0.04)[0],
                                  stated.pairs["USDJPY"].forward_points(0.04)[0], places=6)

    def test_the_callers_clock_beats_the_files_asof_and_the_difference_is_said(self):
        """A valuation in the past against this morning's file is ordinary.

        It is also a fact about what is being priced, so it is reported rather
        than absorbed into a spot date nobody can check.
        """
        f = self.feed("""
            # asof: 2026-08-27
            USDJPY,SPOT,150.25
            USDJPY,2026-09-30,-11.40
        """, today=date(2026, 8, 20))
        self.assertEqual(f.today, date(2026, 8, 20))
        self.assertTrue(any("written as of 2026-08-27" in n for n in f.notes), f.notes)

    def test_tenor_and_dated_rows_mix_in_one_curve(self):
        """A file need not choose: the front dated, the back on tenors."""
        f = self.feed("""
            # asof: 2026-08-27
            USDJPY,SPOT,150.25
            USDJPY,SPOT DATE,2026-08-31
            USDJPY,2026-08-31,-0.40
            USDJPY,2026-09-30,-11.40
            USDJPY,3M,-35.0
            USDJPY,1Y,-146.0
        """)
        self.assertEqual(f.problems, [], f.problems)
        pf = f.pairs["USDJPY"]
        self.assertEqual(pf.tenors, ["2026-09-30", "3M", "1Y"])   # sorted by time
        # The 3M pillar reads back exactly where the curve puts it, which is
        # its own delivery date and not a nominal 30.44-day quarter.
        self.assertAlmostEqual(
            pf.forward_points(pf.times[pf.tenors.index("3M")])[0], -35.0, places=10)
        self.assertAlmostEqual(pf.points_on("2026-09-30")[0], -11.40, places=10)
        self.assertAlmostEqual(pf.points_on("2026-08-30")[0], +0.40, places=10)

    def test_a_label_that_is_neither_says_both(self):
        """The tenor is tried first, so nothing a tenor feed reads moves."""
        f = self.feed("""
            USDJPY,SPOT,150.25
            USDJPY,banana,-11.40
        """)
        self.assertEqual(len(f.problems), 1, f.problems)
        self.assertIn("cannot parse tenor", f.problems[0])
        self.assertIn("not a date either", f.problems[0])

    def test_a_tenor_feed_reads_back_at_its_own_pillars(self):
        """The sample feed's every pillar, and its front and back ends.

        The file states no spot date, but it carries an ``asof`` line, and a
        valuation date is all a spot date needs -- so this feed has one and
        every tenor pillar is placed on its own delivery date.  The pillar
        *values* are untouched by any of that; what moved is where they sit,
        and they read back exactly there.
        """
        f = MarketFeed.load(FEED)
        self.assertEqual(f.problems, [], f.problems)
        pf = f.pairs["USDJPY"]
        self.assertEqual(pf.spot_date, DEFAULT_CALENDARS.spot_date("USDJPY", f.today))
        for t, pts in zip(pf.times, pf.points):
            self.assertAlmostEqual(pf.forward_points(t)[0], pts, places=10)
        self.assertAlmostEqual(pf.forward_points(0.0)[0], 0.0, places=12)
        # the front pillar is still scaled toward zero at the spot date
        half = 0.5 * pf.times[0]
        self.assertAlmostEqual(pf.forward_points(half)[0], 0.5 * pf.points[0], places=10)
        self.assertFalse(pf.forward_points(half)[1])
        self.assertTrue(pf.forward_points(5.0)[1])

    def test_the_feed_is_read_against_the_books_clock(self):
        """``load_for`` is the one caller-facing spelling, and it injects it.

        Every screen and every command goes through it, so a dated feed cannot
        be placed one way on a screen and another way in the batch command
        beside it.
        """
        import tempfile
        from volkit.feed import load_for

        class FakeBook:
            clock = Clock(datetime(2026, 8, 27, 10, 0, tzinfo=UTC))
            calendars = DEFAULT_CALENDARS

        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp, True)
        path = Path(tmp) / "f.csv"
        path.write_text("USDJPY,SPOT,150.25\nUSDJPY,2026-09-30,-11.40\n", encoding="utf-8")
        f = load_for(FakeBook(), path)
        self.assertEqual(f.today, date(2026, 8, 27))
        self.assertEqual(f.pairs["USDJPY"].spot_date, date(2026, 8, 31))


class TestImpliedCrossFeed(unittest.TestCase):
    """A cross the file does not quote is implied from its two legs.

    Two spot rates and two swap points are all an implied cross rate has ever
    been, and a file that quotes EURUSD and USDJPY *is* quoting EURJPY.  The
    arithmetic lives in `feed.compose_level` and nowhere else: `Book`
    contributes only the workbook's opinion about which legs a cross has, so
    there is one place for the triangle's signs to be written (§5 item 1 is
    what a second place costs).
    """

    def feed(self, body, **kw):
        import tempfile
        tmp = tempfile.mkdtemp()
        path = Path(tmp) / "cross.csv"
        path.write_text(textwrap.dedent(body).lstrip(), encoding="utf-8")
        self.addCleanup(shutil.rmtree, tmp, True)
        return MarketFeed.load(path, **kw)

    def test_the_feed_itself_implies_a_cross_it_does_not_quote(self):
        """It used to refuse by name, and only ``Book`` knew the triangle.

        Anything holding a feed and not a book -- the feed status route's own
        quote box among them -- got ``no feed for 'EURJPY'`` off a file that
        was quoting both of its legs.
        """
        f = MarketFeed.load(FEED)
        q = f.quote("EURJPY", 0.25)
        a, b = f.quote("EURUSD", 0.25), f.quote("USDJPY", 0.25)
        self.assertTrue(q["derived"])
        self.assertEqual(q["via"], "EURUSD and USDJPY")
        self.assertAlmostEqual(q["spot"], a["spot"] * b["spot"], places=12)
        self.assertAlmostEqual(q["forward"], a["forward"] * b["forward"], places=12)
        # the points are the cross's own, in the cross's own pips, and are
        # never the legs' points added
        self.assertEqual(q["pip"], 100.0)
        self.assertAlmostEqual(q["points"], (q["forward"] - q["spot"]) * 100.0, places=10)
        self.assertNotAlmostEqual(q["points"], a["points"] + b["points"], places=3)

    def test_a_divided_cross_is_divided(self):
        """EURGBP is EURUSD over GBPUSD, and the sign comes from one place."""
        f = MarketFeed.load(FEED)
        q = f.quote("EURGBP", 0.25)
        a, b = f.quote("EURUSD", 0.25), f.quote("GBPUSD", 0.25)
        self.assertEqual(q["via"], "EURUSD and GBPUSD")
        self.assertAlmostEqual(q["spot"], a["spot"] / b["spot"], places=12)
        self.assertAlmostEqual(q["forward"], a["forward"] / b["forward"], places=12)

    def test_a_cross_nobody_declared_still_has_its_dollar_legs(self):
        """GBPJPY is not in the sample workbook, and the market quotes it.

        The legs of a cross are a fact about its name, not a decision -- which
        is why `cross.dollar_legs` exists and why CONFIG stopped asking for
        them.  Refusing a level for want of a row in a spreadsheet that has
        nothing to do with the feed was the same refusal-by-name in a
        different place.
        """
        f = MarketFeed.load(FEED)
        self.assertNotIn("GBPJPY", f)
        q = f.quote("GBPJPY", 0.25)
        self.assertEqual(q["via"], "GBPUSD and USDJPY")
        self.assertAlmostEqual(q["forward"],
                               f.quote("GBPUSD", 0.25)["forward"]
                               * f.quote("USDJPY", 0.25)["forward"], places=12)

    def test_half_a_triangle_is_still_a_refusal_and_names_the_missing_leg(self):
        """A guessed leg is a level nobody published wearing a published one."""
        f = MarketFeed.load(FEED)
        with self.assertRaises(FeedError) as caught:
            f.quote("GBPNZD", 0.25)
        self.assertIn("NZDUSD", str(caught.exception))
        self.assertIn("GBPUSD and NZDUSD", str(caught.exception))

    def test_the_workbook_wins_when_it_names_a_crosss_legs(self):
        """A sheet that says something is not second-guessed by a convention."""
        f = MarketFeed.load(FEED)
        named = {"EURJPY": ("EURUSD", "USDJPY")}
        q = f.quote("EURJPY", 0.25, lambda p: named.get(p))
        self.assertEqual(q["via"], "EURUSD and USDJPY")
        # and a pair the caller has no opinion about falls through to the
        # dollar legs rather than being refused
        self.assertEqual(f.quote("AUDJPY", 0.25, lambda p: named.get(p))["via"],
                         "AUDUSD and USDJPY")

    def test_the_book_still_gets_exactly_the_numbers_it_did(self):
        """The composition moved into the feed; not one figure moved with it."""
        book = Book.from_excel(BOOK, ASOF)
        book.feed = MarketFeed.load(FEED)
        for pair, via in [("EURJPY", "EURUSD and USDJPY"),
                          ("EURGBP", "EURUSD and GBPUSD"),
                          ("AUDJPY", "AUDUSD and USDJPY")]:
            level = book.market_level(pair, 0.25)
            legs = via.split(" and ")
            a = book.market_level(legs[0], 0.25)
            b = book.market_level(legs[1], 0.25)
            self.assertTrue(level["feed"])
            self.assertTrue(level["derived"])
            self.assertEqual(level["via"], via)
            expect = (a["forward"] * b["forward"] if pair != "EURGBP"
                      else a["forward"] / b["forward"])
            self.assertAlmostEqual(level["forward"], expect, places=12)
        # and a dollar pair the file quotes is untouched by any of it
        self.assertFalse(book.market_level("EURUSD", 0.25)["derived"])

    def test_an_implied_cross_reaches_the_near_side_too(self):
        """The overnight of a cross is the two legs' overnights, composed.

        Read **on the date** rather than at a time, so each leg is placed
        against its own spot date: a cross of a T+1 pair and a T+2 pair has
        two different dates at one ``t``, and the tom-next is a day wide.
        """
        f = self.feed("""
            # asof: 2026-08-27
            EURUSD,SPOT,1.0842
            EURUSD,SPOT DATE,2026-08-31
            EURUSD,2026-08-28,0.03
            EURUSD,2026-08-31,0.12
            EURUSD,2026-09-30,5.8
            USDJPY,SPOT,150.25
            USDJPY,SPOT DATE,2026-08-31
            USDJPY,2026-08-28,-0.06
            USDJPY,2026-08-31,-0.24
            USDJPY,2026-09-30,-11.4
        """)
        self.assertEqual(f.problems, [], f.problems)
        for when in ["2026-08-27", "2026-08-30", "2026-08-31", "2026-09-30"]:
            q = f.quote_on("EURJPY", when)
            a, b = f.quote_on("EURUSD", when), f.quote_on("USDJPY", when)
            self.assertTrue(q["derived"], when)
            self.assertAlmostEqual(q["forward"], a["forward"] * b["forward"],
                                   places=10, msg=when)
        # the cross's points are zero on its own spot date, by definition,
        # and the near side is on the other side of it
        self.assertAlmostEqual(f.quote_on("EURJPY", "2026-08-31")["points"], 0.0, places=10)
        self.assertGreater(f.quote_on("EURJPY", "2026-08-30")["points"], 0.0)
        self.assertLess(f.quote_on("EURJPY", "2026-09-30")["points"], 0.0)

    def test_a_dated_read_of_a_cross_needs_both_legs_to_have_a_spot_date(self):
        """A leg with no valuation date behind it has no dates on it, and says so.

        A file that states a spot date, or one loaded against a clock, has one
        for every pair -- the spot date is derived from the valuation date, and
        that is what lets a *tenor* pillar be placed on its own delivery date.
        With neither there is nothing to derive it from, and a dated read is
        refused by name rather than answered from a guessed origin.
        """
        f = self.feed("""
            EURUSD,SPOT,1.0842
            EURUSD,1M,5.8
            USDJPY,SPOT,150.25
            USDJPY,1M,-11.4
        """)                            # no asof line, and loaded with no clock
        self.assertIsNone(f.today)
        self.assertIsNone(f.pairs["USDJPY"].spot_date)
        with self.assertRaises(FeedError) as caught:
            f.quote_on("EURJPY", "2026-09-30")
        self.assertIn("spot date", str(caught.exception))


class TestFeedRefresh(unittest.TestCase):
    """Picking up spot that has just been published."""

    def service(self, feed):
        from volkit.webapp import BookService
        return BookService(str(BOOK), ASOF, feed_path=str(feed))

    def test_a_rewritten_feed_reads_as_stale_until_it_is_refreshed(self):
        import shutil, tempfile, os, time
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "feed.csv"
            shutil.copy(FEED, path)
            service = self.service(path)
            self.assertFalse(service.feed_state()["stale"])
            text = path.read_text(encoding="utf-8").replace("USDJPY,SPOT,150.25", "USDJPY,SPOT,151.25")
            self.assertNotEqual(text, path.read_text(encoding="utf-8"))
            path.write_text(text, encoding="utf-8")
            os.utime(path, (time.time() + 5, time.time() + 5))
            self.assertTrue(service.feed_state()["stale"])
            r = service.refresh_feed({"legs": [{"pair": "USDJPY", "expiry": "3M"}]})
            self.assertFalse(r["feed"]["stale"])
            self.assertAlmostEqual(r["legs"][0]["spot"], 151.25)

    def test_a_leg_the_feed_cannot_quote_keeps_its_place_with_the_reason(self):
        service = self.service(FEED)
        r = service.refresh_feed({"legs": [{"pair": "USDJPY", "expiry": "1M"},
                                           {"pair": "GBPNZD", "expiry": "1M"},
                                           {"pair": "USDJPY", "expiry": "not a date"}]})
        self.assertEqual(len(r["legs"]), 3)
        self.assertEqual(r["legs"][0]["error"], "")
        # Half a triangle is still a refusal: the file has GBPUSD and no
        # NZDUSD.  The reason names the pair and says both halves were tried.
        self.assertIn("GBPNZD", r["legs"][1]["error"])
        self.assertIn("legs", r["legs"][1]["error"])
        self.assertIsNone(r["legs"][1]["spot"])
        # The expiry is a separate failure and does not take the market with
        # it -- and the row still holds its place either way.
        self.assertTrue(r["legs"][2]["error"])
        self.assertEqual(r["legs"][2]["expiry"], "")
        self.assertEqual([q["index"] for q in r["legs"]], [0, 1, 2])

    def test_a_cross_is_filled_from_the_legs_the_file_does_quote(self):
        """The old bug: this route asked the feed for the pair *by name*.

        The file quotes EURUSD and USDJPY and not EURJPY, so Fill refused the
        cross while the pricing grid underneath it priced the very same leg
        off the very same file -- ``Book.market_level`` composes it.  One
        place reads a level, and this is now that place too.
        """
        service = self.service(FEED)
        r = service.refresh_feed({"legs": [{"pair": "EURJPY", "expiry": "3M"}]})
        q = r["legs"][0]
        self.assertEqual(q["error"], "")
        self.assertTrue(q["derived"])
        self.assertEqual(q["via"], "EURUSD and USDJPY")
        priced = service.price({"legs": [{"pair": "EURJPY", "expiry": "3M",
                                          "strike": "ATM", "type": "C"}]})["legs"][0]
        self.assertAlmostEqual(priced["forward"], q["forward"], places=9)

    def test_the_points_are_the_ones_the_pricer_would_use(self):
        """Filling a leg must not put a different market in front of it."""
        service = self.service(FEED)
        r = service.refresh_feed({"legs": [{"pair": "USDJPY", "expiry": "3M"}]})
        q = r["legs"][0]
        priced = service.price({"legs": [{"pair": "USDJPY", "expiry": "3M", "strike": "ATM",
                                          "type": "C"}]})
        self.assertAlmostEqual(priced["legs"][0]["spot"], q["spot"])
        self.assertAlmostEqual(priced["legs"][0]["forward"], q["forward"], places=9)

    def test_refreshing_without_a_feed_says_so(self):
        from volkit.webapp import BookService
        from volkit.feed import FeedError
        service = BookService(str(BOOK), ASOF)
        with self.assertRaises(FeedError):
            service.refresh_feed({"legs": []})


class TestTheThreeMarketBoxes(unittest.TestCase):
    """The pricing screen shows one box each for spot, the forward and the expiry.

    The forward box is the **outright**, not points over a pip divisor, and
    both level boxes are filled from the feed at the leg's own expiry and are
    then editable.  ``pricing.resolve_legs`` is what fills them, and it is the
    same reading the pricer does -- one place for the calendar and one place
    for the level.
    """

    @classmethod
    def setUpClass(cls):
        from volkit.webapp import BookService
        cls.service = BookService(str(BOOK), ASOF, feed_path=str(FEED))

    def rows(self, *legs):
        return self.service.legs({"legs": list(legs)})["legs"]

    def test_every_spelling_of_a_date_comes_back_as_the_one_standard_date(self):
        """Whatever is typed, the box ends up holding ``YYYY-MM-DD``.

        A desk writes a date half a dozen ways and none of them is worth
        making somebody translate by hand -- ``28May24`` least of all, since
        that is the form this package prints in a leg's own label.
        """
        for text in ("2024-05-28", "28May24", "28May2024", "28 May 24", "28 May 2024",
                     "May 28 2024", "May 28, 2024", "28-May-2024", "28-May-24",
                     "2024/05/28", "5/28/2024", "20240528", "2024.05.28"):
            with self.subTest(text):
                row = self.rows({"pair": "USDJPY", "expiry": text})[0]
                self.assertEqual(row["error"], "")
                self.assertEqual(row["expiry"], "2024-05-28")

    def test_the_box_takes_a_spelled_out_tenor_and_a_year_less_date(self):
        """'1wk' is a tenor and '28 May' is the coming twenty-eighth of May.

        Both used to be refused: the tenor because the unit had to be one
        letter, the date because the year was not optional.  The year comes
        from the book's clock (28-Feb-2024 here), never the machine's.
        """
        week = self.rows({"pair": "USDJPY", "expiry": "1wk"})[0]
        self.assertEqual(week["error"], "")
        self.assertEqual(week["expiry"],
                         self.rows({"pair": "USDJPY", "expiry": "1W"})[0]["expiry"])
        for text in ("28 May", "28May", "May 28"):
            with self.subTest(text):
                row = self.rows({"pair": "USDJPY", "expiry": text})[0]
                self.assertEqual(row["error"], "")
                self.assertEqual(row["expiry"], "2024-05-28")
        # A date already past this year is next year's.
        row = self.rows({"pair": "USDJPY", "expiry": "10 Jan"})[0]
        self.assertEqual(row["error"], "")
        self.assertEqual(row["expiry"], "2025-01-10")

    def test_a_tenor_is_resolved_once_on_the_pair_s_own_calendar(self):
        for tenor in ("1W", "8d", "3M", "2y"):
            with self.subTest(tenor):
                row = self.rows({"pair": "USDJPY", "expiry": tenor})[0]
                self.assertEqual(row["error"], "")
                self.assertTrue(self.service.book.calendars.is_business_day(
                    "USDJPY", date.fromisoformat(row["expiry"])))
        # "8d" is eight days and not the eighth of something: the tenor is
        # tried first, exactly as `resolve_expiry` has always done it.  And a
        # day tenor is eight *business* days, which is what O/N is one of --
        # so it lands eleven or twelve calendar days out, not eight.  Adding
        # calendar days to the spot date instead collapsed the short tenors
        # onto each other: dealt on a Wednesday, "1d" and "2d" both came back
        # Thursday, because the two business days taken off at the end
        # swallowed the weekend the addition had just crossed.
        row = self.rows({"pair": "USDJPY", "expiry": "8d"})[0]
        cal = self.service.book.calendars
        today = self.service.book.clock.now.date()
        self.assertEqual(row["expiry"],
                         cal.add_business_days("USDJPY", today, 8).isoformat())
        self.assertGreater(row["days"], 8)

    def test_the_short_tenors_do_not_collapse_onto_one_date(self):
        """One business day apart, each of them, however the weekend falls."""
        seen = [self.rows({"pair": "USDJPY", "expiry": t})[0]["expiry"]
                for t in ("O/N", "1d", "2d", "3d", "4d")]
        self.assertEqual(seen[0], seen[1])          # O/N is one business day
        self.assertEqual(len(set(seen)), 4)         # and the rest are distinct

    def test_a_leg_carries_its_spot_and_settlement_dates(self):
        """The settlement date is a fact the desk confirms on, not an internal.

        It is the spot lag *after* the expiry, on the pair's own calendar, and
        it is the date the forward beside it is a forward to.
        """
        cal = self.service.book.calendars
        for text in ("1M", "2024-05-28"):
            with self.subTest(text):
                row = self.rows({"pair": "USDJPY", "expiry": text})[0]
                expiry = date.fromisoformat(row["expiry"])
                self.assertEqual(row["settle"],
                                 cal.delivery_from_expiry("USDJPY", expiry).isoformat())
                self.assertEqual(row["spot_date"], cal.spot_date(
                    "USDJPY", self.service.book.clock.now.date()).isoformat())
                self.assertTrue(cal.is_settlement_day("USDJPY",
                                                      date.fromisoformat(row["settle"])))

    def test_the_level_is_the_feed_s_at_that_leg_s_own_expiry(self):
        one, three = self.rows({"pair": "USDJPY", "expiry": "1M"},
                               {"pair": "USDJPY", "expiry": "3M"})
        self.assertTrue(one["feed"] and three["feed"])
        # One spot, two forwards: the points are interpolated at each expiry,
        # which is the whole reason the box is refilled when the expiry moves.
        self.assertAlmostEqual(one["spot"], three["spot"], places=12)
        self.assertNotAlmostEqual(one["forward"], three["forward"], places=6)

    def test_a_row_that_cannot_be_read_keeps_its_place_and_its_reason(self):
        rows = self.rows({"pair": "USDJPY", "expiry": "1M"},
                         {"pair": "USDJPY", "expiry": "not a date"},
                         {"pair": "", "expiry": "1M"})
        self.assertEqual([r["index"] for r in rows], [0, 1, 2])
        self.assertEqual(rows[0]["error"], "")
        self.assertTrue(rows[1]["error"])
        self.assertEqual(rows[1]["expiry"], "")
        self.assertIn("currency pair", rows[2]["error"])

    def test_typing_does_not_re_read_the_feed_file(self):
        """The box is refilled on a keystroke; the file is read on a button.

        Going to disk every time somebody paused in the expiry box would make
        an editor's hesitation a file read, and would pick a republished feed
        up underneath a price being looked at -- which is what the auto-load
        switch exists to make a deliberate choice.
        """
        import shutil, tempfile, os, time
        from volkit.webapp import BookService
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "feed.csv"
            shutil.copy(FEED, path)
            service = BookService(str(BOOK), ASOF, feed_path=str(path))
            was = service.legs({"legs": [{"pair": "USDJPY", "expiry": "3M"}]})["legs"][0]
            path.write_text(path.read_text(encoding="utf-8").replace(
                "USDJPY,SPOT,150.25", "USDJPY,SPOT,151.25"), encoding="utf-8")
            os.utime(path, (time.time() + 5, time.time() + 5))
            still = service.legs({"legs": [{"pair": "USDJPY", "expiry": "3M"}]})["legs"][0]
            self.assertAlmostEqual(still["spot"], was["spot"], places=12)
            now = service.refresh_feed(
                {"legs": [{"pair": "USDJPY", "expiry": "3M"}]})["legs"][0]
            self.assertAlmostEqual(now["spot"], 151.25, places=12)


class TestLegMarketOverrides(unittest.TestCase):
    """Two boxes, either of which may be typed over, and the feed fills the rest.

    The screen sends both, so the ordinary case is that both are typed and
    what is priced is what is on the screen.  The interesting cases are the
    partial ones.
    """

    @classmethod
    def setUpClass(cls):
        cls.book = Book.from_excel(BOOK, ASOF).load_all(["USDJPY"])
        cls.book.feed = MarketFeed.load(FEED)

    def one(self, **kw):
        return price_strip(self.book, [OptionLeg("USDJPY", "3M", "ATM", **kw)])["legs"][0]

    def test_the_feed_fills_both_boxes_when_neither_is_typed(self):
        r = self.one()
        self.assertEqual(r["market_source"], "feed")
        self.assertTrue(r["feed_used"])
        self.assertAlmostEqual(r["spot"], 150.25, places=10)
        self.assertNotAlmostEqual(r["forward"], r["spot"], places=6)

    def test_a_typed_forward_is_the_forward_and_spot_stays_the_feed_s(self):
        r = self.one(forward=155.0)
        self.assertAlmostEqual(r["forward"], 155.0, places=12)
        self.assertAlmostEqual(r["spot"], 150.25, places=10)
        self.assertEqual(r["market_source"], "spot from the feed")

    def test_a_typed_spot_leaves_the_forward_to_the_feed(self):
        """Each box falls back on its own.

        Clearing one of two boxes is an ordinary thing to do to one leg of a
        strip, and it must not need the other cleared as well.  The old
        screen had no forward box at all -- points on top of spot -- so a
        typed spot took the whole market with it.
        """
        fed = self.one()
        r = self.one(spot=160.0)
        self.assertAlmostEqual(r["spot"], 160.0, places=12)
        self.assertAlmostEqual(r["forward"], fed["forward"], places=12)
        self.assertEqual(r["market_source"], "forward from the feed")

    def test_both_boxes_typed_are_priced_exactly_as_they_stand(self):
        r = self.one(spot=160.0, forward=159.4)
        self.assertAlmostEqual(r["spot"], 160.0, places=12)
        self.assertAlmostEqual(r["forward"], 159.4, places=12)
        self.assertEqual(r["market_source"], "typed")
        self.assertFalse(r["feed_used"])

    def test_the_points_spelling_still_says_where_the_forward_is(self):
        """``forward_points`` defaults to None and not to zero.

        Nothing else can tell "said nothing about points" from "said the
        forward is at spot" -- and the two want opposite things from the
        feed.  Defaulting to 0.0, as this did while the screen sent points,
        made every leg the second kind the moment the screen stopped sending
        them.
        """
        r = self.one(spot=160.0, forward_points=0.0, pip=100.0)
        self.assertAlmostEqual(r["forward"], 160.0, places=12)
        self.assertEqual(r["market_source"], "typed")
        r = self.one(spot=160.0, forward_points=-45.0, pip=100.0)
        self.assertAlmostEqual(r["forward"], 160.0 - 0.45, places=10)
        # And a forward given outright wins over both of them.
        r = self.one(spot=160.0, forward=159.0, forward_points=-45.0, pip=100.0)
        self.assertAlmostEqual(r["forward"], 159.0, places=12)

    def test_a_filled_box_the_screen_calls_the_feed_s_reports_as_the_feed_s(self):
        """The bug: every leg read ``typed``, on a screen nobody had typed into.

        The pricing screen fills spot and the outright from the feed and then
        posts what is in the boxes, so by the time a leg is priced a feed
        level and a hand-marked one look identical and the old inference --
        "a box with something in it was somebody's" -- called both of them
        typed.  Only the screen knows, so the screen says: ``spot_source`` /
        ``forward_source`` change no number, they name where the number came
        from.  This is what makes `Refresh spot`, which puts every box back
        on the feed, put the *Market* row back with them.
        """
        fed = self.one()                     # blank boxes: the feed fills both
        r = self.one(spot=fed["spot"], forward=fed["forward"],
                     spot_source="feed", forward_source="feed")
        self.assertEqual(r["market_source"], "feed")
        self.assertTrue(r["feed_used"])
        self.assertAlmostEqual(r["spot"], fed["spot"], places=12)
        self.assertAlmostEqual(r["forward"], fed["forward"], places=12)

    def test_a_level_the_screen_says_is_typed_is_typed_however_it_got_there(self):
        fed = self.one()
        r = self.one(spot=fed["spot"], forward=fed["forward"],
                     spot_source="typed", forward_source="typed")
        self.assertEqual(r["market_source"], "typed")
        self.assertFalse(r["feed_used"])
        # Half and half, each half named: a typed outright over the feed's
        # spot, and the other way round.
        self.assertEqual(self.one(spot=fed["spot"], forward=155.0,
                                  spot_source="feed", forward_source="typed"
                                  )["market_source"], "spot from the feed")
        # A typed spot with the feed's *swap* on top of it: the outright box
        # then holds neither party's outright, and saying "forward from the
        # feed" would claim the file published a level it did not.
        self.assertEqual(self.one(spot=160.0, forward=159.5,
                                  spot_source="typed", forward_source="feed"
                                  )["market_source"], "swap from the feed")

    def test_a_blank_box_is_the_feed_s_whatever_the_caller_says(self):
        """Neither half can be talked out of what actually happened.

        A caller that leaves the box empty and calls it typed still gets the
        feed's level, because that is where it came from; and a caller that
        says nothing at all gets the old inference, which is what the command
        line and any script hold.
        """
        self.assertEqual(self.one(spot_source="typed", forward_source="typed"
                                  )["market_source"], "feed")
        self.assertEqual(self.one()["market_source"], "feed")
        self.assertEqual(self.one(spot=160.0, forward=159.4)["market_source"], "typed")

    def test_a_forward_on_its_own_with_no_feed_is_a_whole_market(self):
        """This model carries no discount curve, so spot has nothing to add.

        The alternative was the old fallback, spot = 1.0, which would price a
        yen option 150 times away from the level in the box beside it.
        """
        book = Book.from_excel(BOOK, ASOF).load_all(["USDJPY"])
        r = price_strip(book, [OptionLeg("USDJPY", "3M", "ATM", forward=155.0)])["legs"][0]
        self.assertTrue(r["ok"], r["error"])
        self.assertAlmostEqual(r["spot"], 155.0, places=12)
        self.assertAlmostEqual(r["forward"], 155.0, places=12)
        self.assertEqual(r["market_source"], "typed")


class TestTimestampReading(unittest.TestCase):
    """One place reads a timestamp, and it reads what the tool writes.

    The listed panel, the events route and the session loader had each
    patched ISO 8601 into shape for themselves, and had done it differently:
    only one of them understood a trailing Z or an offset, so a string that
    parsed on one screen failed on the next.  Worse, that one *dropped* the
    offset and then stamped the result UTC, which reads 19:00+09:00 as 19:00Z.
    """

    def test_iso_8601_is_read_the_way_the_tool_writes_it(self):
        from volkit.timeutil import UTC, parse_datetime
        want = datetime(2024, 2, 28, 12, 0, tzinfo=UTC)
        for text in ("2024-02-28 12:00", "2024-02-28T12:00", "2024-02-28T12:00:00",
                     "2024-02-28T12:00:00Z", "2024-02-28T12:00:00+00:00"):
            with self.subTest(text):
                self.assertEqual(parse_datetime(text), want)
        # The tabular formats are tried first and are untouched.
        self.assertEqual(parse_datetime("2/28/2024 12:00"), want)
        self.assertEqual(parse_datetime("28-Feb-24"),
                         datetime(2024, 2, 28, tzinfo=UTC))
        # And what it prints, it reads.
        self.assertEqual(parse_datetime(want.isoformat()), want)

    def test_an_offset_is_converted_and_not_thrown_away(self):
        from volkit.timeutil import UTC, parse_datetime
        self.assertEqual(parse_datetime("2026-09-11T19:00+09:00"),
                         datetime(2026, 9, 11, 10, 0, tzinfo=UTC))

    def test_nonsense_still_says_so(self):
        from volkit.timeutil import parse_datetime
        with self.assertRaises(ValueError) as caught:
            parse_datetime("next tuesday")
        self.assertIn("next tuesday", str(caught.exception))

    def test_the_listed_panel_takes_its_expiry_straight_from_the_browser(self):
        from volkit.listed import _normalise_expiry
        self.assertEqual(_normalise_expiry(" 2026-09-11T19:00Z "), "2026-09-11T19:00Z")
        with self.assertRaises(ValueError):
            _normalise_expiry("   ")


class TestCrossLevelsFromTheLegs(unittest.TestCase):
    """A cross the feed quotes only through its legs still has a level.

    The bug: the feed publishes EURUSD and USDJPY and therefore publishes
    EURJPY, but every level lookup asked the feed for the pair by name and
    refused.  On the market-maker screen that made a loaded feed invisible --
    a quote written against an absolute strike came back "there is no forward
    feed for EURJPY" while the pricing screen was quoting both its legs off
    the same file.
    """

    def book(self, pairs=("EURJPY", "EURGBP", "GBPNZD")):
        from volkit.feed import MarketFeed
        book = Book.from_excel(BOOK, ASOF).load_all(list(pairs))
        book.feed = MarketFeed.load(FEED)
        return book

    def test_a_cross_is_the_product_of_its_legs(self):
        book = self.book()
        t = 0.25
        level = book.market_level("EURJPY", t)
        self.assertTrue(level["feed"])
        self.assertTrue(level["derived"])
        self.assertEqual(level["via"], "EURUSD and USDJPY")
        legs = (book.feed.quote("EURUSD", t), book.feed.quote("USDJPY", t))
        self.assertAlmostEqual(level["forward"], legs[0]["forward"] * legs[1]["forward"], places=12)
        self.assertAlmostEqual(level["spot"], legs[0]["spot"] * legs[1]["spot"], places=12)
        # The points are the cross's own, in the cross's own pips, and never
        # the legs' points added together.
        self.assertAlmostEqual(level["spot"] + level["points"] / level["pip"],
                               level["forward"], places=12)
        self.assertEqual(level["pip"], 100.0)

    def test_a_cross_of_two_same_side_legs_divides(self):
        """EURGBP is EURUSD / GBPUSD, not EURUSD * GBPUSD."""
        book = self.book()
        t = 0.25
        level = book.market_level("EURGBP", t)
        a, b = book.feed.quote("EURUSD", t), book.feed.quote("GBPUSD", t)
        self.assertAlmostEqual(level["forward"], a["forward"] / b["forward"], places=12)
        self.assertEqual(level["via"], "EURUSD and GBPUSD")

    def test_a_pair_the_feed_quotes_itself_is_not_derived(self):
        book = self.book(["EURUSD"])
        level = book.market_level("EURUSD", 0.25)
        self.assertTrue(level["feed"])
        self.assertFalse(level["derived"])
        self.assertEqual(level["via"], "")

    def test_a_leg_the_feed_does_not_carry_is_still_a_refusal(self):
        """Half a triangle is not a level, and is refused rather than guessed."""
        book = self.book()
        self.assertNotIn("NZDUSD", book.feed.pairs)
        level = book.market_level("GBPNZD", 0.25)
        self.assertFalse(level["feed"])
        self.assertIsNone(level["forward"])
        self.assertIsNone(book.forward_at("GBPNZD", 0.25))

    def test_the_market_maker_prices_an_absolute_strike_on_a_cross(self):
        """The bug, on the screen it was found on."""
        from volkit import marketmaker as mm
        book = self.book(["EURJPY"])
        panel = mm.check_panel_from_request({
            "pair": "EURJPY", "cut": "NY",
            "text": "1M ATM 8.2/8.6\n1M, 162.00, 8.4/8.8\n3M, 162.00, 8.4/8.8\n"})
        rows = panel.run(book)["market"]["rows"]
        self.assertEqual(len(rows), 3)
        for row in rows:
            self.assertIsNotNone(row["model"], row["raw"])
            self.assertEqual([w for w in row["warnings"] if "forward feed" in w], [])

    def test_the_derivation_is_said_once_and_not_once_a_tenor(self):
        from volkit import marketmaker as mm
        book = self.book(["EURJPY"])
        panel = mm.check_panel_from_request({
            "pair": "EURJPY", "cut": "NY",
            "text": "1M, 162.00, 8.4/8.8\n2M, 162.00, 8.4/8.8\n3M, 162.00, 8.4/8.8\n"})
        notes = [n for n in panel.run(book)["market"]["notes"] if "triangle" in n]
        self.assertEqual(len(notes), 1, notes)
        self.assertIn("EURUSD and USDJPY", notes[0])

    def test_every_screen_reads_the_one_lookup(self):
        """``analytics._forward_at`` is the same number as ``market_level``."""
        from volkit.analytics import _forward_at
        book = self.book(["EURJPY"])
        fwd, real, note = _forward_at(book, "EURJPY", 0.25)
        self.assertTrue(real)
        self.assertEqual(fwd, book.market_level("EURJPY", 0.25)["forward"])
        self.assertIn("triangle", note)


if __name__ == "__main__":
    unittest.main()
